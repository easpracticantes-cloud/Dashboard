"""Tests del adaptador Excel Autobits."""

import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

from domain.autobits.fields import suggest_mapping
from infrastructure.autobits.excel_adapter import ExcelAutobitsAdapter


def _build_sample_xlsx() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.append(
        [
            "Proveedor",
            "NIT",
            "Compra",
            "Reserva",
            "No. Factura",
            "Valor",
            "Fecha",
            "Concepto",
        ]
    )
    ws.append(
        [
            "Hotel Andino SAS",
            "900123456",
            "C-1001",
            "R-550",
            "FE-7788",
            850000,
            "2026-08-20",
            "Hospedaje grupo",
        ]
    )
    ws.append(["", "", "", "", "", "", "", ""])
    path = Path(__file__).parent / "_tmp_autobits.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_xlsx():
    path = _build_sample_xlsx()
    yield path
    path.unlink(missing_ok=True)


def test_suggest_mapping_detects_spanish_columns():
    columns = ["Proveedor", "NIT", "Compra", "Reserva", "Valor"]
    mapping = suggest_mapping(columns)
    assert mapping["proveedor"] == "Proveedor"
    assert mapping["nit"] == "NIT"
    assert mapping["numero_compra"] == "Compra"
    assert mapping["valor"] == "Valor"


def test_suggest_mapping_real_autobits_export_columns():
    """Columnas reales del Excel Autobits del usuario."""
    from domain.autobits.fields import AUTOBITS_EXPORT_COLUMNS

    mapping = suggest_mapping(list(AUTOBITS_EXPORT_COLUMNS))
    assert mapping["nit"] == "NIT/CC Proveedor (Orden de Compra)"
    assert mapping["proveedor"] == "Nombre Proveedor (Orden de Compra)"
    assert mapping["numero_compra"] == "Codigo Orden de compra"
    assert mapping["numero_reserva"] == "Codigo Reserva"
    assert mapping["fecha"] == "Fecha de ejecución (Reserva)"
    assert mapping["concepto"] == "Nombre concepto"
    assert mapping["valor"] == "Total"
    assert mapping["observaciones"] == "OBSERVACIONES"
    assert mapping["estado_compra"] == "estado de la compra"
    assert mapping["numero_documento"] == "Codigo Factura proveedor"


def test_excel_adapter_preview_and_parse(sample_xlsx):
    adapter = ExcelAutobitsAdapter()
    preview = adapter.preview(sample_xlsx)
    assert preview.total_rows == 1
    assert "Proveedor" in preview.columns
    assert preview.suggested_mapping["proveedor"] == "Proveedor"

    parsed = adapter.parse(sample_xlsx, preview.suggested_mapping)
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.proveedor == "Hotel Andino SAS"
    assert row.nit == "900123456"
    assert row.numero_compra == "C-1001"
    assert row.valor == 850000.0
    assert parsed.skipped_empty == 0


def test_excel_adapter_ignora_hoja_portada(tmp_path):
    wb = Workbook()
    cover = wb.active
    cover.title = "Portada"
    cover.append(["Autobits Systems"])
    cover.append(["All rights reserved"])
    data = wb.create_sheet("Semana")
    data.append(["Proveedor", "NIT", "Valor", "Fecha"])
    data.append(["Acme", "9001", 50000, "2026-08-18"])
    path = tmp_path / "portada.xlsx"
    wb.save(path)

    preview = ExcelAutobitsAdapter().preview(path)
    assert preview.sheet_name == "Semana"
    assert preview.total_rows == 1


def _real_export_headers():
    return [
        "NIT/CC Proveedor (Orden de Compra)",
        "Nombre Proveedor (Orden de Compra)",
        "Codigo Orden de compra",
        "Referencia (Orden de Compra)",
        "Fecha de compra",
        "Codigo Reserva",
        "Fecha de ejecución (Reserva)",
        "Codigo Factura proveedor",
        "Nombre concepto",
        "Total",
        "OBSERVACIONES",
        "estado de la compra",
    ]


def test_suggest_mapping_no_usa_referencia_como_com():
    from domain.autobits.fields import looks_like_autobits_export

    columns = _real_export_headers()
    mapping = suggest_mapping(columns)
    assert mapping["numero_compra"] == "Codigo Orden de compra"
    assert mapping["numero_documento"] == "Codigo Factura proveedor"
    assert mapping["fecha"] == "Fecha de ejecución (Reserva)"
    assert looks_like_autobits_export(columns)


def test_excel_adapter_export_real_factura_nit_fecha_dinero(tmp_path):
    from datetime import date, datetime

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.append(["Autobits Systems — All rights reserved"])
    ws.append([""])
    ws.append(_real_export_headers())
    ws.append(
        [
            900123456.0,
            "Hotel Andino SAS",
            "COM007441",
            "HIN36005",
            datetime(2026, 8, 18),
            "EAS002722",
            "20/08/2026",
            "FE-6920",
            "Almuerzo",
            "1.234.567,89",
            "pendiente",
            "Cerrada",
        ]
    )
    ws.append(
        [
            "900.123.457-1",
            "TOTAL",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            0,
            "",
            "",
        ]
    )
    path = tmp_path / "autobits_real.xlsx"
    wb.save(path)

    adapter = ExcelAutobitsAdapter()
    preview = adapter.preview(path)
    assert "Codigo Orden de compra" in preview.columns
    assert preview.suggested_mapping["numero_documento"] == "Codigo Factura proveedor"
    assert preview.total_rows >= 1

    parsed = adapter.parse(path, preview.suggested_mapping)
    rows = [r for r in parsed.rows if r.numero_compra]
    assert len(rows) == 1
    row = rows[0]
    assert row.numero_compra == "COM007441"
    assert row.numero_reserva == "EAS002722"
    assert row.numero_documento == "FE-6920"
    assert row.nit == "900123456"
    assert row.proveedor == "Hotel Andino SAS"
    assert row.valor == 1234567.89
    assert row.fecha == date(2026, 8, 20).isoformat()


def test_excel_adapter_dinero_colombiano_y_fechas(tmp_path):
    from datetime import date

    wb = Workbook()
    ws = wb.active
    ws.append(["Proveedor", "NIT", "Valor", "Fecha"])
    ws.append(["Acme", "9001", "$14.300", "20-ago-2026"])
    ws.append(["Beta", 800200300.0, "850000,50", "18/08/26"])
    path = tmp_path / "formatos.xlsx"
    wb.save(path)

    mapping = {
        "proveedor": "Proveedor",
        "nit": "NIT",
        "valor": "Valor",
        "fecha": "Fecha",
    }
    parsed = ExcelAutobitsAdapter().parse(path, mapping)
    by_name = {r.proveedor: r for r in parsed.rows}
    assert by_name["Acme"].valor == 14300.0
    assert by_name["Acme"].fecha == date(2026, 8, 20).isoformat()
    assert by_name["Beta"].nit == "800200300"
    assert by_name["Beta"].valor == 850000.5
    assert by_name["Beta"].fecha == date(2026, 8, 18).isoformat()


def test_excel_adapter_une_hojas_con_mismo_encabezado(tmp_path):
    wb = Workbook()
    a = wb.active
    a.title = "Semana1"
    headers = ["Nombre Proveedor (Orden de Compra)", "Codigo Orden de compra", "Total"]
    a.append(headers)
    a.append(["Hotel A", "COM000001", 1000])
    b = wb.create_sheet("Semana2")
    b.append(headers)
    b.append(["Hotel B", "COM000002", 2000])
    path = tmp_path / "multi.xlsx"
    wb.save(path)

    parsed = ExcelAutobitsAdapter().parse(
        path,
        {
            "proveedor": "Nombre Proveedor (Orden de Compra)",
            "numero_compra": "Codigo Orden de compra",
            "valor": "Total",
        },
    )
    coms = {r.numero_compra for r in parsed.rows}
    assert coms == {"COM000001", "COM000002"}
