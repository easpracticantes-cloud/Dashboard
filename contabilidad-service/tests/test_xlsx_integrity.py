"""El XLSX de Cruce debe ser ZIP/XML válido y reabrirse sin reparación."""

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

from config.settings import get_settings  # noqa: E402
from domain.cruce.export_row import CruceExportRow  # noqa: E402
from domain.enums import DocumentStatus  # noqa: E402
from infrastructure.cruce.workbook_builder import CruceWorkbookBuilder, MASTER_TEMPLATE_PATH  # noqa: E402
from infrastructure.cruce.xlsx_integrity import validate_xlsx_bytes  # noqa: E402
from infrastructure.persistence.database import SessionLocal, init_db  # noqa: E402
from infrastructure.persistence.models import DocumentModel, ProviderModel  # noqa: E402


def _vaciar():
    from infrastructure.persistence.database import engine
    from sqlalchemy import text

    with engine.begin() as conn:
        conn.execute(text("PRAGMA foreign_keys = OFF"))
        tablas = [
            row[0]
            for row in conn.execute(text("SELECT name FROM sqlite_master WHERE type='table'"))
            if not row[0].startswith("sqlite_")
        ]
        for tabla in tablas:
            conn.execute(text(f'DELETE FROM "{tabla}"'))


@pytest.fixture
def client():
    get_settings.cache_clear()
    init_db()
    _vaciar()
    from api_server import app

    yield TestClient(app)
    _vaciar()


def test_maestro_no_se_modifica_al_generar():
    assert MASTER_TEMPLATE_PATH.exists()
    before = MASTER_TEMPLATE_PATH.read_bytes()
    CruceWorkbookBuilder().build(
        [
            CruceExportRow(
                proveedor="JORGE HERNANDO CASTAÑO GIRALDO",
                nit="70905826-6",
                numero_compra="FPOS-61226",
                factura_cdc="FPOS-61226",
                fecha_ejecucion="2026-08-04",
                valor=14300,
            )
        ],
        year=2026,
    )
    assert MASTER_TEMPLATE_PATH.read_bytes() == before


def test_builder_xlsx_zip_xml_roundtrip_sin_legacy_drawing_roto():
    content = CruceWorkbookBuilder().build(
        [
            CruceExportRow(
                proveedor="JORGE HERNANDO CASTAÑO GIRALDO",
                nit="70905826-6",
                numero_compra="FPOS-61226",
                factura_cdc="FPOS-61226",
                fecha_ejecucion="2026-08-04",
                valor=14300,
            )
        ],
        year=2026,
    )
    validate_xlsx_bytes(content)
    wb = load_workbook(io.BytesIO(content), data_only=False)
    agosto = wb["AGOSTO"]
    dumped = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            dumped.extend(str(v) for v in row if v is not None)
    text = " | ".join(dumped)
    assert "FPOS-61226" in text
    assert "JORGE HERNANDO CASTAÑO GIRALDO" in text
    assert "70905826-6" in text
    assert agosto["B3"].value == "FPOS-61226"
    assert agosto["D3"].value == 14300
    assert "COM005691" not in text
    assert "FV POS" not in text
    from zipfile import ZipFile

    with ZipFile(io.BytesIO(content)) as zf:
        names = zf.namelist()
    assert not any("comment" in n.lower() or n.endswith(".vml") or n.startswith("xl/persons/") for n in names)


def test_export_endpoint_xlsx_integro_factura_sin_cruce(client):
    db = SessionLocal()
    try:
        provider = ProviderModel(nombre="JORGE HERNANDO CASTAÑO GIRALDO", nit="70905826-6")
        doc = DocumentModel(
            filename="fpos-61226.pdf",
            tipo="FACTURA",
            origen="CARGA_MANUAL",
            estado=DocumentStatus.CRUZANDO,
            numero_documento="FPOS-61226",
            fecha_emision="2026-08-04",
            total=14300,
            provider=provider,
        )
        db.add_all([provider, doc])
        db.commit()
        db.refresh(doc)
        doc_id = doc.id
    finally:
        db.close()

    res = client.get(f"/api/cruce-excel/export.xlsx?document_ids={doc_id}")
    assert res.status_code == 200, res.text
    assert res.content[:2] == b"PK"
    validate_xlsx_bytes(res.content)
    wb = load_workbook(io.BytesIO(res.content), data_only=False)
    dumped = " | ".join(
        str(v)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for v in row
        if v is not None
    )
    assert "FPOS-61226" in dumped
    assert "JORGE HERNANDO CASTAÑO GIRALDO" in dumped
    assert "70905826-6" in dumped


def test_export_vacio_sigue_siendo_xlsx_valido(client):
    res = client.get("/api/cruce-excel/export.xlsx")
    assert res.status_code == 200
    validate_xlsx_bytes(res.content)
