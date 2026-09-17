"""Tests API Autobits."""

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

from config.settings import get_settings
from infrastructure.persistence.database import init_db


def _make_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Proveedor", "NIT", "Compra", "Valor", "Fecha"])
    ws.append(["Proveedor Demo", "800111222", "CMP-01", 500000, "2026-08-22"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    get_settings.cache_clear()
    init_db()
    from api_server import app

    return TestClient(app)


def test_autobits_preview_and_import(client):
    xlsx = _make_xlsx_bytes()
    preview_resp = client.post(
        "/api/autobits/preview",
        files={"archivo": ("reporte.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview_resp.status_code == 200
    preview = preview_resp.json()
    assert preview["total_rows"] == 1
    assert preview["preview_id"]
    assert preview["suggested_mapping"]["proveedor"] == "Proveedor"

    import_resp = client.post(
        "/api/autobits/import",
        data={
            "preview_id": preview["preview_id"],
            "mapping_json": __import__("json").dumps(preview["suggested_mapping"]),
            "period_start": "2026-08-16",
            "period_end": "2026-08-22",
        },
    )
    assert import_resp.status_code == 200
    data = import_resp.json()
    assert data["imported_rows"] == 1

    list_resp = client.get("/api/autobits/records")
    assert list_resp.status_code == 200
    assert list_resp.json()["total"] >= 1


def test_autobits_upload_direct_auto_detect(client):
    from infrastructure.ai.excel_ai_analyzer import ExcelAIAnalysis
    from unittest.mock import patch
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws.append(["Proveedor", "NIT", "Compra", "Valor", "Fecha"])
    ws.append(["Hotel IA Demo", "800999111", "CMP-IA-99", 777000, "2026-08-21"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx = buf.getvalue()

    fake = ExcelAIAnalysis(
        mapping={
            "proveedor": "Proveedor",
            "nit": "NIT",
            "numero_compra": "Compra",
            "numero_reserva": None,
            "numero_documento": None,
            "valor": "Valor",
            "fecha": "Fecha",
            "concepto": None,
        },
        period_start="2026-08-16",
        period_end="2026-08-22",
        sheet_notes="Mock IA",
        mode="ia",
    )
    with patch(
        "application.services.autobits_service.ExcelAIAnalyzer.analyze",
        return_value=fake,
    ):
        response = client.post(
            "/api/autobits/upload",
            files={
                "archivo": (
                    "reporte_ia.xlsx",
                    xlsx,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"auto_cruzar": "false"},
        )
    assert response.status_code == 200
    data = response.json()
    assert data["imported_rows"] == 1
    assert data["batch"]["period_start"]
    assert data["batch"]["period_end"]
    assert data["detected_mapping"]
    assert data["analysis_mode"] == "ia"
    assert "proveedor" in data["detected_mapping"]


def test_autobits_rejects_non_excel(client):
    response = client.post(
        "/api/autobits/upload",
        files={"archivo": ("bad.txt", b"hola", "text/plain")},
    )
    assert response.status_code == 400


def test_autobits_upload_binds_folder_and_persists_com(client):
    from unittest.mock import patch

    from infrastructure.ai.excel_ai_analyzer import ExcelAIAnalysis

    folder = client.post("/api/folders", json={"name": "Semana COM persist"}).json()
    folder_id = folder["id"]

    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Codigo Orden de compra",
            "Codigo Factura proveedor",
            "Nombre Proveedor (Orden de Compra)",
            "NIT/CC Proveedor (Orden de Compra)",
            "Total",
        ]
    )
    ws.append(["COM007441", "FE-6920", "Hotel Demo", "9001", 17000])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx = buf.getvalue()

    fake = ExcelAIAnalysis(
        mapping={
            "proveedor": "Nombre Proveedor (Orden de Compra)",
            "nit": "NIT/CC Proveedor (Orden de Compra)",
            "numero_compra": "Codigo Factura proveedor",
            "numero_reserva": None,
            "numero_documento": "Codigo Factura proveedor",
            "valor": "Total",
            "fecha": None,
            "concepto": None,
        },
        period_start="2026-09-13",
        period_end="2026-09-19",
        sheet_notes="Mock IA",
        mode="ia",
    )
    with patch(
        "application.services.autobits_service.ExcelAIAnalyzer.analyze",
        return_value=fake,
    ):
        response = client.post(
            "/api/autobits/upload",
            files={
                "archivo": (
                    "autobits.xlsx",
                    xlsx,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"auto_cruzar": "false", "force": "true", "folder_id": str(folder_id)},
        )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["folder"]
    assert data["folder"]["id"] == folder_id
    assert data["folder"]["autobits_batch_id"] == data["batch"]["id"]
    assert data["folder_error"] is None
    assert data["records"][0]["numero_compra"] == "COM007441"

    got = client.get(f"/api/folders/{folder_id}")
    assert got.status_code == 200
    assert got.json()["autobits_batch_id"] == data["batch"]["id"]

    records = client.get(f"/api/autobits/records?batch_id={data['batch']['id']}")
    assert records.status_code == 200
    assert records.json()["items"][0]["numero_compra"] == "COM007441"
