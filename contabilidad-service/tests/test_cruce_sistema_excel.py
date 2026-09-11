"""Cruce de Cuentas desde SIG: sin Excel de entrada, Excel de salida estándar."""

import io
import sys
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

from config.settings import get_settings  # noqa: E402
from domain.cruce.export_row import CruceExportRow  # noqa: E402
from domain.cruce.workbook_spec import (  # noqa: E402
    BOSQUE_HEADERS,
    DUSTER_HEADERS,
    LUGER_HEADERS,
    PERIOD_BLOCK_HEADERS,
    standard_sheet_names,
)
from infrastructure.ai.excel_ai_analyzer import ExcelAIAnalysis  # noqa: E402
from infrastructure.cruce.workbook_builder import CruceWorkbookBuilder  # noqa: E402
from infrastructure.persistence.database import init_db  # noqa: E402


def _fake_analysis():
    return ExcelAIAnalysis(
        mapping={
            "proveedor": "Nombre Proveedor (Orden de Compra)",
            "nit": "NIT/CC Proveedor (Orden de Compra)",
            "numero_compra": "Codigo Orden de compra",
            "numero_reserva": "Codigo Reserva",
            "numero_documento": None,
            "valor": "Total",
            "fecha": "Fecha de ejecución (Reserva)",
            "concepto": "Nombre concepto",
            "observaciones": "OBSERVACIONES",
            "estado_compra": "estado de la compra",
        },
        period_start="2026-08-16",
        period_end="2026-08-22",
        mode="ia",
        sheet_notes="test",
    )


def _vaciar_tablas():
    from infrastructure.persistence.database import engine
    from sqlalchemy import text

    with engine.begin() as conn:
        conn.execute(text("PRAGMA foreign_keys = OFF"))
        tablas = [
            row[0]
            for row in conn.execute(text("SELECT name FROM sqlite_master WHERE type='table'"))
            if not row[0].startswith("sqlite_")
        ]
        for tabla in tablas:
            conn.execute(text(f'DELETE FROM "{tabla}"'))


@pytest.fixture
def client():
    get_settings.cache_clear()
    init_db()
    _vaciar_tablas()
    from api_server import app

    yield TestClient(app)
    _vaciar_tablas()


def _autobits_xlsx(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "NIT/CC Proveedor (Orden de Compra)",
            "Nombre Proveedor (Orden de Compra)",
            "Codigo Orden de compra",
            "Codigo Reserva",
            "Fecha de ejecución (Reserva)",
            "estado de la compra",
            "Nombre concepto",
            "Moneda",
            "Total",
            "SI",
            "NO",
            "OBSERVACIONES",
        ]
    )
    for nit, proveedor, compra, reserva, fecha, valor, obs in filas:
        ws.append(
            [nit, proveedor, compra, reserva, fecha, "Activa", "Servicio", "COP", valor, "", "", obs]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _subir_autobits(client, filas):
    with patch(
        "application.services.autobits_service.ExcelAIAnalyzer.analyze",
        return_value=_fake_analysis(),
    ):
        return client.post(
            "/api/autobits/upload",
            files={
                "archivo": (
                    "autobits.xlsx",
                    _autobits_xlsx(filas),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"auto_cruzar": "true"},
        )


def test_analizar_sin_excel_de_cruce(client):
    up = _subir_autobits(
        client,
        [["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-08-20", 150000, ""]],
    )
    assert up.status_code == 200, up.text

    res = client.post("/api/cruce-excel/analizar")
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["archivo"] == "sistema"
    assert data["origen"] == "SIG"
    assert data["lectura"]["filas_leidas"] >= 1
    assert data["batch"]["imported_rows"] == 1
    comparacion = data["comparacion"]
    assert comparacion
    assert comparacion[0]["lado_autobits"]["compra"] == "COM001"
    assert "FACTURA/CDC" in comparacion[0]["faltas"] or comparacion[0]["faltas"]


def test_export_xlsx_estructura_estandar(client):
    _subir_autobits(
        client,
        [
            ["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-03-10", 150000, ""],
            ["800222", "VENTAS DUSTER", "COM002", "EAS002", "2026-04-01", 80000, ""],
        ],
    )
    client.post("/api/cruce-excel/analizar")
    export = client.get("/api/cruce-excel/export.xlsx")
    assert export.status_code == 200, export.text
    assert "spreadsheet" in export.headers["content-type"]
    assert "Cruce_Cuentas_" in export.headers.get("content-disposition", "")
    assert export.content[:2] == b"PK"
    assert len(export.content) > 200
    alias = client.get("/api/cruce-excel/export")
    assert alias.status_code == 200, alias.text
    assert alias.content[:2] == b"PK"

    wb = load_workbook(io.BytesIO(export.content))
    year = 2026
    assert tuple(wb.sheetnames) == standard_sheet_names(year)
    enero = wb[standard_sheet_names(year)[0]]
    headers = [(enero.cell(2, c).value or "").strip() for c in range(1, 7)]
    assert tuple(headers) == PERIOD_BLOCK_HEADERS
    dumped = _xlsx_text(export.content)
    assert "COM001" not in dumped
    assert "COM002" not in dumped
    assert "Hotel Demo SAS" not in dumped
    duster = wb["VENTAS_DUSTER"]
    assert [duster.cell(2, c).value for c in range(1, 11)] == list(DUSTER_HEADERS)
    assert duster["C3"].value in (None, "")
    assert duster["J3"].value is None
    bosque = wb["CDC BOSQUE DE PALMAS"]
    assert [bosque.cell(9, c).value for c in range(1, 11)] == list(BOSQUE_HEADERS)
    luger = wb["PRECOMPRA LUGER 2026"]
    assert [luger.cell(5, c).value for c in range(2, 7)] == list(LUGER_HEADERS)


def test_workbook_builder_no_inventa_ceros():
    builder = CruceWorkbookBuilder()
    content = builder.build(
        [
            CruceExportRow(
                proveedor="Restaurante Demo",
                numero_compra="COM100",
                numero_reserva="EAS100",
                fecha_ejecucion="2026-06-01",
                valor=None,
                factura_cdc=None,
                fecha_pago=None,
            )
        ],
        year=2026,
    )
    wb = load_workbook(io.BytesIO(content))
    mayo = wb["MAYO - JULIO"]
    assert mayo["D3"].value is None
    assert mayo["E3"].value is None
    assert mayo["F3"].value is None


def test_workbook_builder_formula_sum_y_especiales():
    builder = CruceWorkbookBuilder()
    content = builder.build(
        [
            CruceExportRow(
                proveedor="VENTAS DUSTER",
                nit="900",
                numero_compra="COM9",
                numero_reserva="EAS9",
                fecha_ejecucion="2026-02-01",
                valor=Decimal("1000"),
                concepto="Tour",
            ),
            CruceExportRow(
                proveedor="Restaurante X",
                numero_compra="COM8",
                fecha_ejecucion="2026-02-02",
                valor=Decimal("2000"),
            ),
        ],
        year=2026,
    )
    wb = load_workbook(io.BytesIO(content))
    enero = wb["AÑO  2026 ENERO - ABRIL"]
    assert any(
        isinstance(cell.value, str) and "SUMIF" in cell.value
        for row in enero.iter_rows(min_row=1, max_row=20, max_col=10)
        for cell in row
    )
    duster = wb["VENTAS_DUSTER"]
    assert duster["I3"].value == 1000
    assert duster["I4"].value == "=SUM(I3:I3)"


def test_matching_ambiguo_no_es_exacto():
    from domain.enums import MatchType
    from domain.matching.matching_engine import MatchCandidate, MatchingEngine

    engine = MatchingEngine()
    best = MatchCandidate(autobits_record_id=1, score=80, match_type=MatchType.MATCH_EXACTO)
    second = MatchCandidate(autobits_record_id=2, score=78, match_type=MatchType.MATCH_PROBABLE)
    engine._flag_ambiguity(best, [best, second])
    assert "ambiguo" in best.reasons
    assert best.match_type == MatchType.MATCH_PROBABLE


def test_workbook_no_trunca_filas_largas():
    builder = CruceWorkbookBuilder()
    rows = [
        CruceExportRow(
            proveedor="Restaurante Largo",
            numero_compra=f"COM{i:04d}",
            fecha_ejecucion="2026-02-01",
            valor=Decimal("1000"),
            fecha_pago="2026-02-10" if i % 2 == 0 else None,
        )
        for i in range(90)
    ]
    wb = load_workbook(io.BytesIO(builder.build(rows, year=2026)))
    enero = wb["AÑO  2026 ENERO - ABRIL"]
    assert enero["B3"].value == "COM0000"
    assert enero["B92"].value == "COM0089"
    assert "SUMIF" in str(enero["D93"].value)


def test_raw_json_llena_columnas_especiales_si_existen():
    builder = CruceWorkbookBuilder()
    content = builder.build(
        [
            CruceExportRow(
                proveedor="VENTAS DUSTER",
                nit="900",
                numero_compra="COM9",
                referencia_oc="REF-OC",
                numero_reserva="EAS9",
                fecha_ejecucion="2026-02-01",
                valor=Decimal("1000"),
                precio_terceros=Decimal("800"),
                concepto="Tour",
            )
        ],
        year=2026,
    )
    duster = load_workbook(io.BytesIO(content))["VENTAS_DUSTER"]
    assert duster["D3"].value == "REF-OC"
    assert duster["J3"].value == 800


def test_analizar_no_depende_del_excel_historico(client):
    _subir_autobits(
        client,
        [["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-08-20", 150000, ""]],
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "AGOSTO"
    ws.append(["Hotel Demo SAS"])
    ws.append(
        [
            "FECHA DE EJECUCIÓN",
            "ORDEN DE COMPRA",
            "REF.",
            "VALOR",
            "FACTURA/CDC",
            "FECHA DE PAGO",
        ]
    )
    ws.append(["2026-08-20", "COM999", "EAS999", 1, "FV HIST", "2026-08-21"])
    buf = io.BytesIO()
    wb.save(buf)
    client.post(
        "/api/cruce-excel/upload",
        files={
            "archivo": (
                "historico.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"aplicar": "true"},
    )
    res = client.post("/api/cruce-excel/analizar")
    assert res.status_code == 200, res.text
    compras = [r["lado_autobits"].get("compra") for r in res.json()["comparacion"]]
    assert "COM001" in compras


def test_pendientes_ignora_snapshot_historico(client):
    _subir_autobits(
        client,
        [["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-08-20", 150000, ""]],
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "AGOSTO"
    ws.append(["Hotel Demo SAS"])
    ws.append(list(PERIOD_BLOCK_HEADERS))
    ws.append(["2026-08-20", "COM999", "EAS999", 1, "FV HIST", "2026-08-21"])
    buf = io.BytesIO()
    wb.save(buf)
    client.post(
        "/api/cruce-excel/upload",
        files={
            "archivo": (
                "historico.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"aplicar": "true"},
    )
    res = client.get("/api/cruce-excel/pendientes")
    assert res.status_code == 200, res.text
    compras = [r["lado_autobits"].get("compra") for r in res.json()["comparacion"]]
    assert "COM001" in compras
    faltan = res.json()["pendientes"]["por_tipo"].get("FALTA_EN_CRUCE") or []
    assert not faltan


def test_fecha_en_periodo_no_inventa_semana():
    from application.services.cruce_excel_service import CruceExcelService

    assert CruceExcelService._fecha_en_periodo("2026-08-20", "2026-08-16", "2026-08-22")
    assert CruceExcelService._fecha_en_periodo("2026-08-20T15:00:00", "2026-08-16", "2026-08-22")
    assert not CruceExcelService._fecha_en_periodo(None, "2026-08-16", "2026-08-22")
    assert not CruceExcelService._fecha_en_periodo("2025-01-01", "2026-08-16", "2026-08-22")


def test_to_money_or_none_no_inventa_cero_en_texto():
    from domain.utils.money import to_money_or_none

    assert to_money_or_none(None) is None
    assert to_money_or_none("") is None
    assert to_money_or_none("N/A") is None
    assert to_money_or_none("150000") == Decimal("150000.00")


def test_upload_historico_sigue_existiendo(client):
    """El endpoint de upload no se elimina (compatibilidad); el flujo de producto no lo usa."""
    _subir_autobits(
        client,
        [["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-08-20", 150000, ""]],
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "AGOSTO"
    ws.append(["Hotel Demo SAS"])
    ws.append(
        [
            "FECHA DE EJECUCIÓN",
            "ORDEN DE COMPRA",
            "REF.",
            "VALOR",
            "FACTURA/CDC",
            "FECHA DE PAGO",
        ]
    )
    ws.append(["2026-08-20", "COM001", "EAS001", 150000, "FV 1", "2026-08-21"])
    buf = io.BytesIO()
    wb.save(buf)
    res = client.post(
        "/api/cruce-excel/upload",
        files={
            "archivo": (
                "historico.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"aplicar": "true"},
    )
    assert res.status_code == 200, res.text


def test_export_usa_plantilla_maestra_sin_datos_historicos():
    builder = CruceWorkbookBuilder()
    content = builder.build(
        [
            CruceExportRow(
                proveedor="Hotel Nuevo SAS",
                numero_compra="COM-NUEVO-1",
                numero_reserva="EAS-NUEVO",
                fecha_ejecucion="2026-03-15",
                valor=Decimal("250000"),
                factura_cdc="FE-4589",
            )
        ],
        year=2026,
    )
    wb = load_workbook(io.BytesIO(content))
    assert tuple(wb.sheetnames) == standard_sheet_names(2026)
    dumped = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            dumped.extend(str(v) for v in row if v is not None)
    joined = " | ".join(dumped)
    assert "COM-NUEVO-1" in joined
    assert "FE-4589" in joined
    assert "COM005691" not in joined
    assert "FV POS" not in joined
    enero = wb["AÑO  2026 ENERO - ABRIL"]
    assert tuple(enero.cell(2, c).value for c in range(1, 7)) == PERIOD_BLOCK_HEADERS


def _crear_factura(
    *,
    nombre: str,
    nit: str,
    numero: str,
    fecha: str,
    total: float,
    filename: str,
):
    from domain.enums import DocumentStatus
    from infrastructure.persistence.database import SessionLocal
    from infrastructure.persistence.models import DocumentModel, ProviderModel

    db = SessionLocal()
    try:
        provider = ProviderModel(nombre=nombre, nit=nit)
        doc = DocumentModel(
            filename=filename,
            tipo="FACTURA",
            origen="CARGA_MANUAL",
            estado=DocumentStatus.PROCESADO,
            numero_documento=numero,
            total=total,
            fecha_emision=fecha,
            provider=provider,
        )
        db.add_all([provider, doc])
        db.commit()
        db.refresh(doc)
        return doc.id
    finally:
        db.close()


def test_export_document_ids_no_mezcla_facturas(client):
    _subir_autobits(
        client,
        [["900111", "Hotel Demo SAS", "COM001", "EAS001", "2026-03-10", 150000, ""]],
    )
    id_a = _crear_factura(
        nombre="EMPRESA XYZ",
        nit="900123456-7",
        numero="FE-AAAA",
        fecha="2026-03-11",
        total=1250000,
        filename="factura-a.pdf",
    )
    id_b = _crear_factura(
        nombre="EMPRESA XYZ",
        nit="900123456-7",
        numero="FE-BBBB",
        fecha="2026-03-12",
        total=800000,
        filename="factura-b.pdf",
    )

    excel_a = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_a}")
    assert excel_a.status_code == 200, excel_a.text
    text_a = _xlsx_text(excel_a.content)
    assert "FE-AAAA" in text_a
    assert "FE-BBBB" not in text_a
    assert "COM001" not in text_a

    excel_b = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_b}")
    assert excel_b.status_code == 200, excel_b.text
    text_b = _xlsx_text(excel_b.content)
    assert "FE-BBBB" in text_b
    assert "FE-AAAA" not in text_b

    excel_a2 = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_a}")
    assert "FE-AAAA" in _xlsx_text(excel_a2.content)
    assert "FE-BBBB" not in _xlsx_text(excel_a2.content)


def test_export_solo_empresas_con_factura_y_datos_de_la_factura(client):
    _subir_autobits(
        client,
        [
            ["900111", "Empresa A Autobits", "COM-A", "EAS-A", "2026-01-01", 111, ""],
            ["900222", "Empresa B", "COM-B", "EAS-B", "2026-01-02", 222, ""],
            ["900333", "Empresa C", "COM-C", "EAS-C", "2026-01-03", 333, ""],
            ["900444", "Empresa D Autobits", "COM-D", "EAS-D", "2026-01-04", 444, ""],
        ],
    )
    id_a = _crear_factura(
        nombre="Empresa A",
        nit="900111",
        numero="FAC-00125",
        fecha="2026-03-10",
        total=1500000,
        filename="factura-a.pdf",
    )
    id_d = _crear_factura(
        nombre="Empresa D",
        nit="900444",
        numero="FAC-00999",
        fecha="2026-03-11",
        total=750000,
        filename="factura-d.pdf",
    )

    solo_a = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_a}")
    assert solo_a.status_code == 200, solo_a.text
    text_a = _xlsx_text(solo_a.content)
    assert "Empresa A" in text_a
    assert "FAC-00125" in text_a
    assert "1500000" in text_a or "1,500,000" in text_a
    assert "Empresa B" not in text_a
    assert "Empresa C" not in text_a
    assert "Empresa D" not in text_a
    assert "COM-A" not in text_a
    assert "COM-B" not in text_a
    wb_a = load_workbook(io.BytesIO(solo_a.content))
    enero = wb_a["AÑO  2026 ENERO - ABRIL"]
    assert enero["A1"].value == "Empresa A"
    assert enero["A3"].value.year == 2026
    assert enero["A3"].value.month == 3
    assert enero["A3"].value.day == 10
    assert enero["B3"].value == "FAC-00125"
    assert enero["D3"].value == 1500000
    assert enero["E3"].value == "FAC-00125"

    ambas = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_a},{id_d}")
    text_ad = _xlsx_text(ambas.content)
    assert "Empresa A" in text_ad
    assert "Empresa D" in text_ad
    assert "FAC-00125" in text_ad
    assert "FAC-00999" in text_ad
    assert "Empresa B" not in text_ad
    assert "Empresa C" not in text_ad
    assert "COM-B" not in text_ad
    assert "COM-C" not in text_ad


def test_export_batch_id_sin_document_ids_no_vuelca_autobits(client):
    up = _subir_autobits(
        client,
        [
            ["900111", "Empresa A", "COM005691", "EAS-A", "2026-03-10", 111, ""],
            ["900222", "Empresa B", "COM-B", "EAS-B", "2026-03-11", 222, ""],
        ],
    )
    assert up.status_code == 200, up.text
    batch_id = up.json()["batch"]["id"]
    client.post(f"/api/cruce-excel/analizar?batch_id={batch_id}")
    export = client.get(f"/api/cruce-excel/export.xlsx?batch_id={batch_id}")
    assert export.status_code == 200, export.text
    dumped = _xlsx_text(export.content)
    assert "COM005691" not in dumped
    assert "Empresa A" not in dumped
    assert "Empresa B" not in dumped
    assert "COM-B" not in dumped


def test_export_batch_id_no_reconstruye_document_ids_aunque_haya_facturas(client):
    """Aunque el lote tenga facturas cruzadas, batch_id solo no las mete al Excel."""
    up = _subir_autobits(
        client,
        [["900111", "Empresa A Autobits", "COM-A", "EAS-A", "2026-03-10", 111, ""]],
    )
    assert up.status_code == 200, up.text
    batch_id = up.json()["batch"]["id"]
    doc_id = _crear_factura(
        nombre="Empresa A",
        nit="900111",
        numero="FAC-00125",
        fecha="2026-03-10",
        total=1500000,
        filename="factura-a.pdf",
    )
    client.get(f"/api/cruce-excel/export.xlsx?document_ids={doc_id}")
    solo_lote = client.get(f"/api/cruce-excel/export.xlsx?batch_id={batch_id}")
    assert solo_lote.status_code == 200, solo_lote.text
    dumped = _xlsx_text(solo_lote.content)
    assert "FAC-00125" not in dumped
    assert "Empresa A" not in dumped
    assert "COM-A" not in dumped


def test_export_no_sustituye_factura_por_com_de_autobits(client):
    _subir_autobits(
        client,
        [["900111", "Empresa A Autobits", "COM005691", "EAS-HIST", "2026-01-01", 999, ""]],
    )
    doc_id = _crear_factura(
        nombre="Empresa A",
        nit="900111",
        numero="FAC-00125",
        fecha="2026-03-10",
        total=1500000,
        filename="factura-fac.pdf",
    )
    export = client.get(f"/api/cruce-excel/export.xlsx?document_ids={doc_id}")
    assert export.status_code == 200, export.text
    dumped = _xlsx_text(export.content)
    assert "FAC-00125" in dumped
    assert "COM005691" not in dumped
    assert "2026-01-01" not in dumped
    wb = load_workbook(io.BytesIO(export.content))
    enero = wb["AÑO  2026 ENERO - ABRIL"]
    assert enero["B3"].value == "FAC-00125"
    assert enero["D3"].value == 1500000
    assert enero["E3"].value == "FAC-00125"


def test_export_cruce_de_otra_factura_no_se_mezcla(client):
    from domain.enums import CrossingStatus
    from infrastructure.persistence.database import SessionLocal
    from infrastructure.persistence.models import AccountCrossingModel

    id_a = _crear_factura(
        nombre="Empresa A",
        nit="900111",
        numero="FE-AAAA",
        fecha="2026-03-10",
        total=1000000,
        filename="a.pdf",
    )
    id_b = _crear_factura(
        nombre="Empresa B",
        nit="900222",
        numero="FE-BBBB",
        fecha="2026-03-11",
        total=2000000,
        filename="b.pdf",
    )
    db = SessionLocal()
    try:
        db.add_all(
            [
                AccountCrossingModel(
                    document_id=id_a,
                    estado=CrossingStatus.PENDIENTE,
                    fecha_pago="2026-03-20",
                    proveedor_nombre="Cruce A",
                ),
                AccountCrossingModel(
                    document_id=id_b,
                    estado=CrossingStatus.PENDIENTE,
                    fecha_pago="2026-03-28",
                    proveedor_nombre="Cruce B",
                ),
            ]
        )
        db.commit()
    finally:
        db.close()

    excel_a = client.get(f"/api/cruce-excel/export.xlsx?document_ids={id_a}")
    assert excel_a.status_code == 200, excel_a.text
    dumped = _xlsx_text(excel_a.content)
    assert "FE-AAAA" in dumped
    assert "FE-BBBB" not in dumped
    assert "Empresa B" not in dumped
    assert "Cruce B" not in dumped
    wb = load_workbook(io.BytesIO(excel_a.content))
    enero = wb["AÑO  2026 ENERO - ABRIL"]
    assert enero["A1"].value == "Empresa A"
    assert enero["D3"].value == 1000000
    pago = enero["F3"].value
    assert "2026-03-28" not in str(pago)
    if pago is not None:
        assert getattr(pago, "day", 20) == 20 or "2026-03-20" in str(pago)


def test_export_bloquea_facturas_en_proceso(client):
    from domain.enums import DocumentStatus
    from infrastructure.persistence.database import SessionLocal
    from infrastructure.persistence.models import DocumentModel, ProviderModel

    db = SessionLocal()
    try:
        provider = ProviderModel(nombre="Hotel Proceso", nit="900999")
        doc = DocumentModel(
            filename="factura-proc.pdf",
            tipo="FACTURA",
            origen="CARGA_MANUAL",
            estado=DocumentStatus.PROCESANDO,
            numero_documento="FE-PROC",
            provider=provider,
        )
        db.add_all([provider, doc])
        db.commit()
        db.refresh(doc)
        doc_id = doc.id
    finally:
        db.close()

    res = client.get(f"/api/cruce-excel/export.xlsx?document_ids={doc_id}")
    assert res.status_code == 409


def _xlsx_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    values = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values.extend(str(v) for v in row if v is not None)
    return " | ".join(values)
