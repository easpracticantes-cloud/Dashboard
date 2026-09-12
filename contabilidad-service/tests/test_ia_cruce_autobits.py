"""Claude relaciona facturas con Autobits; el Excel usa ese vínculo real."""

import io
import sys
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

import pytest
from fastapi.testclient import TestClient

from application.services.crossing_service import CrossingService  # noqa: E402
from application.services.document_processing_service import DocumentProcessingService  # noqa: E402
from config.settings import get_settings  # noqa: E402
from domain.enums import DocumentStatus  # noqa: E402
from infrastructure.persistence.database import SessionLocal, init_db  # noqa: E402
from infrastructure.persistence.models import (  # noqa: E402
    AutobitsRecordModel,
    DocumentModel,
    ImportBatchModel,
    ProviderModel,
)


@pytest.fixture
def client():
    get_settings.cache_clear()
    init_db()
    from api_server import app

    return TestClient(app)


class _FakeAI:
    def __init__(self, payload):
        self.payload = payload

    def verify(self):
        return True

    def extract_invoice(self, *_a, **_k):
        raise AssertionError("no extract_invoice")

    def extract_custom(self, *_a, **_k):
        raise AssertionError("no extract_custom")

    def extract_json(self, *_a, **_k):
        from infrastructure.ai.ollama_provider import AIExtractionResult

        return AIExtractionResult(ok=True, data=self.payload)


def test_relacionar_con_autobits_solo_acepta_ids_reales():
    doc = SimpleNamespace(
        id=11,
        numero_documento="FPOS-61226",
        fecha_emision="2026-08-04",
        total=14300,
        concepto=None,
        extracted_json=None,
        provider=SimpleNamespace(nombre="JORGE HERNANDO CASTAÑO GIRALDO", nit="70905826-6"),
    )
    rec = SimpleNamespace(
        id=22,
        proveedor="JORGE HERNANDO CASTAÑO GIRALDO",
        nit="70905826-6",
        fecha="2026-08-04",
        valor=14300,
        numero_compra="COM-61226",
        numero_reserva="EAS-1",
        concepto="hospedaje",
    )
    ai = _FakeAI(
        {
            "vinculos": [
                {
                    "document_id": 11,
                    "autobits_record_id": 22,
                    "confianza": 0.95,
                    "razones": ["mismo total", "misma fecha"],
                },
                {
                    "document_id": 99,
                    "autobits_record_id": 22,
                    "confianza": 0.99,
                    "razones": ["id inventado"],
                },
                {
                    "document_id": 11,
                    "autobits_record_id": 22,
                    "confianza": 0.4,
                    "razones": ["baja confianza"],
                },
            ]
        }
    )
    out = DocumentProcessingService(ai=ai).relacionar_con_autobits([doc], [rec])
    assert out["ok"] is True
    assert out["vinculos"] == [
        {
            "document_id": 11,
            "autobits_record_id": 22,
            "confianza": 0.95,
            "razones": ["mismo total", "misma fecha"],
        }
    ]


def test_aplicar_vinculos_ia_y_excel_usa_com_de_esa_factura(client):
    db = SessionLocal()
    try:
        batch = ImportBatchModel(filename="semana.xlsx", imported_rows=1)
        provider = ProviderModel(nombre="JORGE HERNANDO CASTAÑO GIRALDO", nit="70905826-6")
        doc = DocumentModel(
            filename="fpos-ia.pdf",
            tipo="FACTURA",
            origen="CARGA_MANUAL",
            estado=DocumentStatus.CRUZANDO,
            numero_documento="FPOS-61226",
            fecha_emision="2026-08-04",
            total=14300,
            provider=provider,
        )
        rec = AutobitsRecordModel(
            import_batch=batch,
            row_number=1,
            proveedor="JORGE HERNANDO CASTAÑO GIRALDO",
            nit="70905826-6",
            fecha="2026-08-04",
            valor=14300,
            numero_compra="COM-61226",
            numero_reserva="EAS-1",
        )
        db.add_all([batch, provider, doc, rec])
        db.commit()
        db.refresh(doc)
        db.refresh(rec)
        doc_id = doc.id
        CrossingService(db).aplicar_vinculos_ia(
            [
                {
                    "document_id": doc_id,
                    "autobits_record_id": rec.id,
                    "confianza": 0.92,
                    "razones": ["mismo total", "fecha"],
                }
            ],
            usuario="TEST",
        )
        db.commit()
    finally:
        db.close()

    export = client.get(f"/api/cruce-excel/export.xlsx?document_ids={doc_id}")
    assert export.status_code == 200, export.text
    dumped = []
    wb = load_workbook(io.BytesIO(export.content))
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            dumped.extend(str(v) for v in row if v is not None)
    text = " | ".join(dumped)
    assert "FPOS-61226" in text
    assert "COM-61226" in text
    assert "EAS-1" in text
