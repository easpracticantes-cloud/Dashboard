"""Validación e higiene OOXML: el XLSX que se entrega, no el workbook en memoria."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import posixpath
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

_REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_PKG_ID = "Id"
_NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"

_HAZARD_PART = re.compile(
    r"(^xl/comments(/|$))"
    r"|(^xl/comments\d)"
    r"|(comment\d*\.xml$)"
    r"|(^xl/persons/)"
    r"|(threadedComments)"
    r"|(\.vml$)"
    r"|(vmlDrawing)"
    r"|(commentsDrawing)"
    r"|(^xl/drawings/)"
    r"|(^xl/tables/)"
    r"|(^xl/media/)",
    re.I,
)
_HAZARD_REL_TYPE = re.compile(
    r"(comments|vmlDrawing|/drawing|table|/hyperlink|/image|/chart|ctrlProp|slicer|persons)",
    re.I,
)
_LEGACY_DRAWING = re.compile(r"<legacyDrawing\b[^>]*?(?:/>|>.*?</legacyDrawing>)", re.I | re.S)
_DRAWING_EL = re.compile(r"<drawing\b[^>]*?(?:/>|>.*?</drawing>)", re.I | re.S)
_TABLE_PARTS = re.compile(r"<tableParts\b[^>]*>.*?</tableParts>", re.I | re.S)
_TABLE_PART = re.compile(r"<tablePart\b[^>]*?(?:/>|>.*?</tablePart>)", re.I | re.S)
_HYPERLINKS = re.compile(r"<hyperlinks\b[^>]*>.*?</hyperlinks>", re.I | re.S)
_LEGACY_HF = re.compile(r"<legacyDrawingHF\b[^>]*?(?:/>|>.*?</legacyDrawingHF>)", re.I | re.S)
_EMPTY_T_N = re.compile(r'(<c\b[^>]*?)\s+t="n"([^>]*?/>)')
_FILTER_NAME = re.compile(
    r"<definedName\b[^>]*_FilterDatabase[^>]*>.*?</definedName>",
    re.I | re.S,
)


class XlsxIntegrityError(ValueError):
    """El workbook no es un XLSX válido para Microsoft Excel."""


def strip_ooxml_hazards(content: bytes) -> bytes:
    """Quita comentarios/VML/drawings/tablas/hipervínculos del ZIP ya guardado.

    openpyxl puede reescribir esas partes al hacer save() aunque el modelo
    en memoria se haya vaciado. La fuente de verdad es el paquete OOXML.
    """
    if not content or content[:2] != b"PK":
        raise XlsxIntegrityError("El archivo no es un ZIP/XLSX (falta firma PK).")

    with ZipFile(BytesIO(content)) as src:
        kept: list[tuple[str, bytes]] = []
        for info in src.infolist():
            name = info.filename
            if name.endswith("/") or _HAZARD_PART.search(name):
                continue
            data = src.read(name)
            if name.endswith(".xml") or name.endswith(".rels"):
                data = _clean_xml_part(name, data)
            kept.append((name, data))

    names = {name for name, _ in kept}
    cleaned: list[tuple[str, bytes]] = []
    for name, data in kept:
        if name.endswith(".rels"):
            data = _drop_missing_rel_targets(name, data, names)
        elif name == "[Content_Types].xml":
            data = _drop_missing_content_types(data, names)
        cleaned.append((name, data))

    return _write_xlsx_zip(cleaned)


def validate_xlsx_bytes(content: bytes) -> None:
    if not content or content[:2] != b"PK":
        raise XlsxIntegrityError("El archivo no es un ZIP/XLSX (falta firma PK).")
    with ZipFile(BytesIO(content)) as zf:
        broken = zf.testzip()
        if broken:
            raise XlsxIntegrityError(f"ZIP corrupto: {broken}")
        names = set(zf.namelist())
        required = {"[Content_Types].xml", "xl/workbook.xml"}
        missing = required - names
        if missing:
            raise XlsxIntegrityError(f"Faltan partes OOXML: {sorted(missing)}")
        hazards = [n for n in names if _HAZARD_PART.search(n)]
        if hazards:
            raise XlsxIntegrityError(
                "El XLSX aún contiene partes que Excel repara: " + ", ".join(sorted(hazards))
            )
        for name in names:
            if name.endswith("/") or name.endswith((".bin", ".emf", ".png", ".jpeg", ".jpg")):
                continue
            if name.endswith(".xml") or name.endswith(".rels"):
                raw = zf.read(name)
                try:
                    ET.fromstring(raw)
                except ET.ParseError as exc:
                    raise XlsxIntegrityError(f"XML inválido en {name}: {exc}") from exc
                text = raw.decode("utf-8", errors="replace")
                if "legacyDrawing" in text:
                    raise XlsxIntegrityError(f"{name} todavía contiene <legacyDrawing>.")
        if "xl/styles.xml" in names:
            wbrels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")
            if "relationships/styles" not in wbrels:
                raise XlsxIntegrityError(
                    "workbook.xml.rels perdió el vínculo a styles.xml (Excel lo marca inválido)."
                )
        _assert_package_relationships(zf, names)

    wb = load_workbook(BytesIO(content), data_only=False)
    if not wb.sheetnames:
        raise XlsxIntegrityError("El workbook no tiene hojas.")
    buf = BytesIO()
    wb.save(buf)
    roundtrip = buf.getvalue()
    load_workbook(BytesIO(roundtrip), data_only=False)
    with ZipFile(BytesIO(roundtrip)) as zf:
        leftovers = [n for n in zf.namelist() if _HAZARD_PART.search(n)]
        if leftovers:
            raise XlsxIntegrityError(
                "El round-trip de openpyxl reintrodujo partes peligrosas: "
                + ", ".join(sorted(leftovers))
            )


def _clean_xml_part(name: str, data: bytes) -> bytes:
    text = data.decode("utf-8")
    if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
        text = _LEGACY_DRAWING.sub("", text)
        text = _DRAWING_EL.sub("", text)
        text = _TABLE_PARTS.sub("", text)
        text = _TABLE_PART.sub("", text)
        text = _HYPERLINKS.sub("", text)
        text = _LEGACY_HF.sub("", text)
        text = _EMPTY_T_N.sub(r"\1\2", text)
    elif name.endswith(".rels"):
        text = _strip_hazard_relationships(text)
    elif name == "[Content_Types].xml":
        text = _strip_hazard_content_types(text)
    elif name == "xl/workbook.xml":
        text = _FILTER_NAME.sub("", text)
    return text.encode("utf-8")


def _strip_hazard_relationships(text: str) -> str:
    def drop(match: re.Match[str]) -> str:
        blob = match.group(0)
        if _HAZARD_REL_TYPE.search(blob):
            return ""
        return blob

    return re.sub(r"<Relationship\b[^>]*?/>", drop, text)


def _strip_hazard_content_types(text: str) -> str:
    def drop(match: re.Match[str]) -> str:
        blob = match.group(0)
        if re.search(
            r"comments|vml|person|drawing|table|threaded",
            blob,
            re.I,
        ):
            return ""
        return blob

    text = re.sub(r"<Override\b[^>]*?/>", drop, text)
    text = re.sub(r'<Default\b[^>]*Extension="vml"[^>]*?/>', "", text)
    return text


def _relationships_base_dir(rel_path: str) -> str:
    """Directorio de la parte dueña del .rels (no el de la carpeta _rels).

    En OOXML, Target relativo se resuelve desde la parte fuente
    (p. ej. xl/workbook.xml), no desde xl/_rels/.
    """
    rel_path = rel_path.replace("\\", "/")
    marker = "/_rels/"
    if marker in rel_path:
        return rel_path.split(marker, 1)[0]
    if rel_path.startswith("_rels/"):
        return ""
    return posixpath.dirname(rel_path)


def _rel_target_exists(rel_path: str, target: str, names: set[str]) -> bool:
    if target.startswith("http://") or target.startswith("https://") or target.startswith("mailto:"):
        return True
    base = _relationships_base_dir(rel_path)
    if target.startswith("/"):
        resolved = target.lstrip("/")
    else:
        resolved = posixpath.normpath(posixpath.join(base, target) if base else target)
    if resolved.startswith("../"):
        resolved = posixpath.normpath(resolved)
    return resolved in names or resolved + "/" in names


def _drop_missing_rel_targets(rel_path: str, data: bytes, names: set[str]) -> bytes:
    text = data.decode("utf-8")

    def drop(match: re.Match[str]) -> str:
        blob = match.group(0)
        target_m = re.search(r'Target="([^"]+)"', blob)
        if not target_m:
            return blob
        if not _rel_target_exists(rel_path, target_m.group(1), names):
            return ""
        return blob

    return re.sub(r"<Relationship\b[^>]*?/>", drop, text).encode("utf-8")


def _drop_missing_content_types(data: bytes, names: set[str]) -> bytes:
    text = data.decode("utf-8")

    def drop(match: re.Match[str]) -> str:
        blob = match.group(0)
        part = re.search(r'PartName="([^"]+)"', blob)
        if not part:
            return blob
        part_name = part.group(1).lstrip("/")
        if part_name not in names:
            return ""
        return blob

    return re.sub(r"<Override\b[^>]*?/>", drop, text).encode("utf-8")


def _assert_package_relationships(zf: ZipFile, names: set[str]) -> None:
    for name in names:
        if not name.endswith(".rels"):
            continue
        root = ET.fromstring(zf.read(name))
        for rel in root:
            if rel.tag.split("}")[-1] != "Relationship":
                continue
            rid = rel.attrib.get(_PKG_ID) or rel.attrib.get("Id")
            target = rel.attrib.get("Target") or ""
            rel_type = rel.attrib.get("Type") or ""
            if _HAZARD_REL_TYPE.search(rel_type):
                raise XlsxIntegrityError(
                    f"{name} todavía declara relationship {rel_type} id={rid!r}."
                )
            if target and not _rel_target_exists(name, target, names):
                raise XlsxIntegrityError(
                    f"{name} apunta a Target inexistente {target!r} (id={rid!r})."
                )

    for name in names:
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(zf.read(name))
        rels_name = name.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
        rel_ids: set[str] = set()
        if rels_name in names:
            rels = ET.fromstring(zf.read(rels_name))
            for rel in rels:
                rid = rel.attrib.get(_PKG_ID) or rel.attrib.get("Id")
                if rid:
                    rel_ids.add(rid)
        for el in root.iter():
            tag = el.tag.split("}")[-1]
            rid = el.attrib.get(_REL_ID) or el.attrib.get("id")
            if tag in {"legacyDrawing", "drawing", "tablePart"}:
                raise XlsxIntegrityError(f"{name} todavía contiene <{tag}>.")
            if rid and rid not in rel_ids:
                raise XlsxIntegrityError(
                    f"{name} usa r:id={rid!r} que no existe en {rels_name}."
                )

    if "[Content_Types].xml" in names:
        root = ET.fromstring(zf.read("[Content_Types].xml"))
        for el in root:
            if el.tag.split("}")[-1] != "Override":
                continue
            part = (el.attrib.get("PartName") or "").lstrip("/")
            if part and part not in names:
                raise XlsxIntegrityError(
                    f"[Content_Types].xml declara Override huérfano: {part}"
                )


def _write_xlsx_zip(parts: list[tuple[str, bytes]]) -> bytes:
    order = {
        "[Content_Types].xml": 0,
        "_rels/.rels": 1,
        "xl/workbook.xml": 2,
        "xl/_rels/workbook.xml.rels": 3,
    }
    parts = sorted(parts, key=lambda item: (order.get(item[0], 50), item[0]))
    buf = BytesIO()
    with ZipFile(buf, "w", compression=ZIP_DEFLATED) as dst:
        for name, data in parts:
            dst.writestr(name, data)
    return buf.getvalue()


def maestro_sha256(path: Path) -> str:
    import hashlib

    return hashlib.sha256(path.read_bytes()).hexdigest()
