"""Genera el Excel de salida: una sola hoja tabular Cruce de cuentas."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from domain.cruce.export_row import CruceExportRow
from domain.cruce.workbook_spec import (
    FONT_NAME,
    HEADER_FILL_BLUE,
    SINGLE_SHEET_HEADERS,
    SINGLE_SHEET_NAME,
)
from infrastructure.cruce.xlsx_integrity import (
    XlsxIntegrityError,
    strip_ooxml_hazards,
    validate_xlsx_bytes,
)

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
_MONEY_FMT = "#,##0"
_DATE_FMT = "YYYY-MM-DD"

MASTER_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "templates" / "CRUCE_DE_CUENTAS_MAESTRO.xlsx"
)


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


def _as_date(value: str | None):
    if not value:
        return None
    text = str(value)[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return value


def _as_number(value: Decimal | None):
    if value is None:
        return None
    return float(value)


class CruceWorkbookBuilder:
    """Export = 1 hoja con las columnas del estándar (sin copiar el libro histórico)."""

    def build(self, rows: list[CruceExportRow], *, year: int | None = None) -> bytes:
        _ = year or self._infer_year(rows)
        wb = Workbook()
        ws = wb.active
        ws.title = SINGLE_SHEET_NAME
        self._write_single_sheet(ws, rows)

        tmp_path = None
        try:
            fd, tmp_path = tempfile.mkstemp(prefix="Cruce_Cuentas_", suffix=".tmp.xlsx")
            os.close(fd)
            wb.save(tmp_path)
            wb.close()
            content = strip_ooxml_hazards(Path(tmp_path).read_bytes())
            validate_xlsx_bytes(content)
            return content
        except XlsxIntegrityError:
            raise
        finally:
            if tmp_path:
                Path(tmp_path).unlink(missing_ok=True)

    def _infer_year(self, rows: list[CruceExportRow]) -> int:
        years = [r.year() for r in rows if r.year()]
        if years:
            return max(years)
        return date.today().year

    def _write_single_sheet(self, ws: Worksheet, rows: list[CruceExportRow]) -> None:
        fill = _fill(HEADER_FILL_BLUE)
        align = Alignment(horizontal="center", wrap_text=True, vertical="center")
        for col, header in enumerate(SINGLE_SHEET_HEADERS, start=1):
            cell = ws.cell(1, col, header)
            cell.fill = fill
            cell.font = _HEADER_FONT
            cell.alignment = align
            cell.border = _THIN

        ordered = sorted(
            rows,
            key=lambda r: (
                (r.proveedor or "").casefold(),
                r.fecha_ejecucion or "",
                r.numero_compra or "",
                r.document_id or 0,
            ),
        )
        for i, row in enumerate(ordered, start=2):
            # Preferir COM del contramarcado; nunca rellenar OC con el nº de factura
            oc_com = row.contramarcado_com or row.numero_compra
            values = [
                row.proveedor,
                row.nit,
                _as_date(row.fecha_ejecucion),
                oc_com,
                row.numero_reserva,
                _as_number(row.valor),
                row.factura_cdc,
                _as_date(row.fecha_pago),
                row.contramarcado,
                row.contramarcado_status,
                row.contramarcado_source,
                row.concepto,
                row.estado_compra,
                row.match_type,
                row.observaciones,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(i, col, value)
                cell.border = _THIN
                if col in (3, 8) and isinstance(value, date):
                    cell.number_format = _DATE_FMT
                if col == 6 and value is not None:
                    cell.number_format = _MONEY_FMT

        widths = (28, 14, 14, 18, 14, 12, 16, 14, 42, 12, 12, 22, 14, 12, 28)
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(SINGLE_SHEET_HEADERS))
        ws.auto_filter.ref = f"A1:{last_col}{max(1, 1 + len(ordered))}"
