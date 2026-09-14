"""Valida el Excel generado: una sola hoja tabular Cruce de cuentas."""
from __future__ import annotations

import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from domain.cruce.export_row import CruceExportRow
from domain.cruce.workbook_spec import SINGLE_SHEET_HEADERS, SINGLE_SHEET_NAME
from infrastructure.cruce.workbook_builder import CruceWorkbookBuilder

OUT = ROOT / "dataset" / "demo" / "Cruce_Cuentas_validacion_estructura.xlsx"


def main() -> int:
    sample = [
        CruceExportRow(
            proveedor="RESTAURANTE DEMO",
            numero_compra="COM0001",
            numero_reserva="EAS0001",
            fecha_ejecucion="2026-02-10",
            valor=Decimal("89000"),
            factura_cdc="FV POS 1",
            fecha_pago="2026-02-15",
        ),
        CruceExportRow(
            proveedor="VENTAS DUSTER",
            nit="900000",
            numero_compra="COM0002",
            numero_reserva="EAS0002",
            fecha_ejecucion="2026-03-01",
            valor=Decimal("120000"),
            concepto="Servicio",
            estado_compra="Activa",
        ),
    ]
    content = CruceWorkbookBuilder().build(sample, year=2026)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(content)

    gen = load_workbook(OUT)
    assert gen.sheetnames == [SINGLE_SHEET_NAME], gen.sheetnames
    ws = gen[SINGLE_SHEET_NAME]
    headers = [ws.cell(1, c).value for c in range(1, 9)]
    assert tuple(headers) == SINGLE_SHEET_HEADERS
    assert ws["D2"].value in {"COM0001", "COM0002"}
    print("OK", OUT, "hoja unica", SINGLE_SHEET_NAME)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
