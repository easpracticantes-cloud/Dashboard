"""Reconstruye CRUCE_DE_CUENTAS_MAESTRO.xlsx desde un Excel operativo.

Solo conserva organización (hojas, anchos, encabezados/etiquetas de tablas
especiales). NO copia filas históricas (COM*/FV POS/proveedores de ejemplo).

Uso:
  py -3 scripts/rebuild_maestro_estructura.py "C:/Users/.../CRUCE DE CUENTAS 2026.xlsx"
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from domain.cruce.fields import fold  # noqa: E402
from infrastructure.cruce.workbook_builder import MASTER_TEMPLATE_PATH  # noqa: E402

_PERIOD_NAMES = {"mayo - julio", "agosto"}
_KEEP_FRAGMENTS = (
    "fecha",
    "orden",
    "ref.",
    "ref ",
    "valor",
    "factura",
    "pago",
    "total",
    "mes",
    "nit",
    "codigo",
    "reserva",
    "estado",
    "description",
    "precio",
    "nombre",
    "comprador",
    "vendedor",
    "cantidad",
    "ingreso",
    "egreso",
    "unidades",
    "saldo",
    "entradas",
    "tour",
    "cdc bosque",
    "ventas_duster",
    "precompra",
    "disponible",
)


def _is_period_sheet(name: str) -> bool:
    folded = fold(name)
    return folded.startswith("ano ") or folded in _PERIOD_NAMES


def _keep_structural(value) -> bool:
    if not isinstance(value, str):
        return False
    text = value.strip()
    if not text or text.startswith("="):
        return False
    folded = fold(text)
    return any(token in folded for token in _KEEP_FRAGMENTS)


def _sanitize_workbook(wb) -> None:
    for key in list(wb.defined_names):
        try:
            if "_FilterDatabase" in str(key):
                del wb.defined_names[key]
        except Exception:
            continue
    for ws in wb.worksheets:
        ws._comments = []
        ws.legacy_drawing = None
        if hasattr(ws, "_hyperlinks"):
            ws._hyperlinks = []
        if hasattr(ws, "_images"):
            ws._images = []
        if hasattr(ws, "_charts"):
            ws._charts = []
        if hasattr(ws, "_drawing"):
            ws._drawing = None
        tables = getattr(ws, "tables", None)
        if tables:
            for table_name in list(tables):
                del tables[table_name]
        if getattr(ws, "auto_filter", None) is not None:
            try:
                ws.auto_filter.ref = None
            except Exception:
                pass
        for row in ws.iter_rows():
            for cell in row:
                if getattr(cell, "comment", None):
                    cell.comment = None
                if getattr(cell, "hyperlink", None):
                    cell.hyperlink = None


def rebuild(source: Path, dest: Path = MASTER_TEMPLATE_PATH) -> Path:
    wb = load_workbook(source, data_only=False)
    _sanitize_workbook(wb)

    for name in wb.sheetnames:
        ws = wb[name]
        for rng in list(ws.merged_cells.ranges):
            try:
                ws.unmerge_cells(str(rng))
            except Exception:
                pass

        if _is_period_sheet(name):
            for row in ws.iter_rows(
                min_row=1,
                max_row=ws.max_row or 1,
                max_col=ws.max_column or 1,
            ):
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    cell.value = None
            continue

        for row in ws.iter_rows(
            min_row=1,
            max_row=ws.max_row or 1,
            max_col=ws.max_column or 1,
        ):
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                if _keep_structural(cell.value):
                    # Quita sufijos de año en títulos estructurales genéricos.
                    if isinstance(cell.value, str) and re.search(
                        r"unidades disponibles", fold(cell.value)
                    ):
                        cell.value = "UNIDADES DISPONIBLES AÑO"
                    continue
                cell.value = None

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    return dest


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    source = Path(argv[1]).expanduser().resolve()
    if not source.exists():
        print(f"No existe: {source}")
        return 1
    out = rebuild(source)
    print(f"OK maestro estructura -> {out} ({out.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
