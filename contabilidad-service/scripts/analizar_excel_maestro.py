"""Inspección estructural del Excel maestro CRUCE DE CUENTAS (solo lectura)."""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def cell_info(cell) -> dict:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        kind = "formula"
    elif value is None:
        kind = "empty"
    else:
        kind = type(value).__name__
    fill = None
    if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb and cell.fill.fgColor.rgb != "00000000":
        fill = str(cell.fill.fgColor.rgb)
    return {
        "addr": cell.coordinate,
        "value": value if not isinstance(value, bytes) else "<bytes>",
        "kind": kind,
        "num_fmt": cell.number_format,
        "font": cell.font.name if cell.font else None,
        "font_size": cell.font.size if cell.font else None,
        "bold": bool(cell.font.bold) if cell.font else False,
        "fill": fill,
    }


def analyze(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    report: dict = {
        "archivo": str(path),
        "hojas": [],
        "nombres": list(wb.sheetnames),
        "defined_names": [n for n in wb.defined_names],
    }
    for name in wb.sheetnames:
        ws = wb[name]
        used = ws.calculate_dimension()
        max_r, max_c = ws.max_row or 0, ws.max_column or 0
        formulas = []
        headers = []
        non_empty = 0
        kinds: Counter[str] = Counter()
        for row in ws.iter_rows(min_row=1, max_row=min(max_r, 200), max_col=min(max_c, 80)):
            for cell in row:
                if cell.value is None:
                    continue
                non_empty += 1
                info = cell_info(cell)
                kinds[info["kind"]] += 1
                if info["kind"] == "formula":
                    formulas.append({"addr": info["addr"], "formula": info["value"]})
                text = str(cell.value).strip().upper()
                if any(
                    token in text
                    for token in (
                        "FECHA",
                        "ORDEN",
                        "FACTURA",
                        "VALOR",
                        "REF",
                        "NIT",
                        "PRECIO",
                        "TOTAL",
                        "MES",
                        "INGRESO",
                        "EGRESO",
                    )
                ):
                    headers.append(info)
        freeze = ws.freeze_panes
        merges = [str(m) for m in ws.merged_cells.ranges]
        widths = {
            get_column_letter(i): ws.column_dimensions[get_column_letter(i)].width
            for i in range(1, min(max_c, 40) + 1)
            if ws.column_dimensions[get_column_letter(i)].width
        }
        heights = {
            r: ws.row_dimensions[r].height
            for r in range(1, min(max_r, 20) + 1)
            if ws.row_dimensions[r].height
        }
        auto_filter = str(ws.auto_filter.ref) if ws.auto_filter and ws.auto_filter.ref else None
        tables = [t.name for t in ws.tables.values()] if hasattr(ws, "tables") else []
        preview = []
        for r in range(1, min(12, max_r) + 1):
            preview.append(
                [ws.cell(r, c).value for c in range(1, min(18, max_c) + 1)]
            )
        report["hojas"].append(
            {
                "nombre": name,
                "used": used,
                "max_row": max_r,
                "max_col": max_c,
                "non_empty_sample": non_empty,
                "kinds": dict(kinds),
                "freeze_panes": freeze,
                "merges": merges[:80],
                "merge_count": len(merges),
                "column_widths": widths,
                "row_heights": heights,
                "auto_filter": auto_filter,
                "tables": tables,
                "sheet_state": ws.sheet_state,
                "formula_count": len(formulas),
                "formulas_sample": formulas[:40],
                "header_like": headers[:40],
                "preview_12x18": preview,
            }
        )
    wb.close()
    return report


if __name__ == "__main__":
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "/data/master.xlsx")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "/data/master_analysis.json")
    data = analyze(src)
    out.write_text(json.dumps(data, ensure_ascii=False, default=str, indent=2), encoding="utf-8")
    print(json.dumps({"hojas": data["nombres"], "detalle": [
        {
            "nombre": h["nombre"],
            "max_row": h["max_row"],
            "max_col": h["max_col"],
            "formulas": h["formula_count"],
            "merges": h["merge_count"],
            "freeze": h["freeze_panes"],
            "filter": h["auto_filter"],
            "tables": h["tables"],
        }
        for h in data["hojas"]
    ]}, ensure_ascii=False, indent=2))
